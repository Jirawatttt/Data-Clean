#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
quarterly_sales_pipeline.py
===========================
Pipeline สรุปรายงาน "SALES DASHBOARD" รายเดือนจาก POS (ไฟล์ .xlsx) -- เลือกโหมดได้

โหมด (--mode, เลือกได้หลายอัน หรือ all)
---------------------------------------
quarter  รวมเป็นรายไตรมาส: ยอดขายตาม "หมวดหมู่" และ "ช่วงเวลา" + ยอดรวม + % จากยอดรวม   (ค่าเริ่มต้น)
weekday  สรุปตามวันในสัปดาห์ (จันทร์-อาทิตย์): ยอดรวม, % จากยอดรวม, เฉลี่ยต่อวันที่เปิดขาย
week     สรุปรายสัปดาห์ (จันทร์-อาทิตย์): ยอดรวม, % จากยอดรวม, เฉลี่ยต่อวันที่เปิดขาย
day      สรุปรายวัน: ยอดรวม, % จากยอดรวม

หมายเหตุ: ไฟล์ต้นทางมีตารางหมวดหมู่/ช่วงเวลาเป็น "ยอดรวมต่อเดือน" เท่านั้น (ไม่ได้แยกรายวัน)
          จึงแบ่งตามวัน/สัปดาห์ได้เฉพาะ "ยอดขาย" จากตารางรายวัน ส่วนหมวดหมู่/ช่วงเวลาใช้ในโหมด quarter

หลักการทำงาน
------------
1. อ่านทุกไฟล์/ทุกชีตใน input แล้ว "หาตาราง" จากชื่อหัวคอลัมน์ (ไม่ฟิกตำแหน่งเซลล์)
2. ตรวจเดือน/ปีจากหัวรายงาน (เช่น "ประจำเดือน สิงหาคม 2569") -> ถ้าไม่เจอใช้ชื่อไฟล์/ชีต
3. ชื่อ field ไม่ตรงกันได้: จับคู่ผ่าน alias (ภาษาไทย/อังกฤษ), เพิ่ม alias เองได้ด้วย --aliases
   คอลัมน์ตัวเลขที่ไม่รู้จักจะถูกรวมให้, คอลัมน์ข้อความที่ไม่รู้จักจะถูกข้าม,
   field/ตาราง/หมวดที่เดือนใดไม่มีนับเป็น 0 และบันทึกไว้ในชีต "ข้อสังเกต"
4. ตรวจความสอดคล้องของยอด (หัวรายงาน vs รายวัน vs หมวดหมู่ vs ช่วงเวลา) -> ชีต "สรุป" + "ข้อสังเกต"
5. ไม่แก้ไฟล์ต้นฉบับ สร้างไฟล์ผลลัพธ์ใหม่เสมอ (ค่าจริง ไม่ใช่สูตร)

วิธีใช้
-------
    pip install -r requirements.txt

    python quarterly_sales_pipeline.py -i ./data                              # รายไตรมาส (default)
    python quarterly_sales_pipeline.py -i Data_Jul.xlsx Data_Aug.xlsx -m weekday
    python quarterly_sales_pipeline.py -i ./data -m week day --quarter 2569-Q3
    python quarterly_sales_pipeline.py -i ./data -m all --csv
    python quarterly_sales_pipeline.py -i ./data --fiscal-start-month 10      # ปีงบประมาณเริ่ม ต.ค.
    python quarterly_sales_pipeline.py -i ./data --aliases my_aliases.json

ตัวอย่าง my_aliases.json (ทั้งสองคีย์ไม่จำเป็นต้องมีครบ):
    {"columns": {"ยอดรวมสุทธิ": "sales", "Dept": "category"},
     "tables":  {"Dept": "category"}}
    ชื่อมาตรฐานของคอลัมน์: date weekday category slot product sales net_sales orders qty discount pct_source
    ชนิดตาราง: daily category slot product

เรียกใช้จากโค้ดอื่น:
    from quarterly_sales_pipeline import run_pipeline
    run_pipeline(["./data"], "./output", modes=["quarter", "weekday"])
"""
from __future__ import annotations

import argparse
import difflib
import json
import logging
import re
import sys
import unicodedata
import warnings
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

log = logging.getLogger("quarter_pipeline")

# ----------------------------------------------------------------------------
# 1) ค่าคงที่ / ตั้งค่า (แก้ตรงนี้เมื่อรูปแบบไฟล์เปลี่ยน หรือใช้ --aliases)
# ----------------------------------------------------------------------------
THAI_MONTHS = {
    1: ["มกราคม", "ม.ค."], 2: ["กุมภาพันธ์", "ก.พ."], 3: ["มีนาคม", "มี.ค."],
    4: ["เมษายน", "เม.ย."], 5: ["พฤษภาคม", "พ.ค."], 6: ["มิถุนายน", "มิ.ย."],
    7: ["กรกฎาคม", "ก.ค."], 8: ["สิงหาคม", "ส.ค."], 9: ["กันยายน", "ก.ย."],
    10: ["ตุลาคม", "ต.ค."], 11: ["พฤศจิกายน", "พ.ย."], 12: ["ธันวาคม", "ธ.ค."],
}
EN_MONTHS = {
    1: ["january", "jan"], 2: ["february", "feb"], 3: ["march", "mar"],
    4: ["april", "apr"], 5: ["may"], 6: ["june", "jun"], 7: ["july", "jul"],
    8: ["august", "aug"], 9: ["september", "sep", "sept"], 10: ["october", "oct"],
    11: ["november", "nov"], 12: ["december", "dec"],
}
MONTH_ABBR_TH = {m: names[1] for m, names in THAI_MONTHS.items()}
MONTH_FULL_TH = {m: names[0] for m, names in THAI_MONTHS.items()}

THAI_WEEKDAYS = ["จันทร์", "อังคาร", "พุธ", "พฤหัส", "ศุกร์", "เสาร์", "อาทิตย์"]
_WEEKDAY_SOURCES = {
    0: ["จันทร์", "จ", "mon", "monday"], 1: ["อังคาร", "อ", "tue", "tuesday"],
    2: ["พุธ", "พ", "wed", "wednesday"], 3: ["พฤหัส", "พฤหัสบดี", "พฤ", "thu", "thursday"],
    4: ["ศุกร์", "ศ", "fri", "friday"], 5: ["เสาร์", "ส", "sat", "saturday"],
    6: ["อาทิตย์", "อา", "sun", "sunday"],
}
WEEKDAY_ALIASES = {t.casefold(): i for i, ts in _WEEKDAY_SOURCES.items() for t in ts}

# ชื่อหัวคอลัมน์ (หลัง casefold + ตัดช่องว่าง) -> ชื่อมาตรฐาน
COLUMN_ALIASES = {
    "วันที่": "date", "date": "date",
    "วัน": "weekday", "weekday": "weekday",
    "หมวดหมู่": "category", "หมวด": "category", "หมวดสินค้า": "category",
    "ประเภทสินค้า": "category", "category": "category",
    "ช่วงเวลา": "slot", "ช่วง": "slot", "เวลา": "slot", "timeslot": "slot", "time": "slot",
    "สินค้า": "product", "รายการสินค้า": "product", "product": "product", "item": "product",
    "ยอดขาย": "sales", "ยอดขายรวม": "sales", "ยอดรวม": "sales", "ยอด": "sales",
    "sales": "sales", "amount": "sales", "revenue": "sales", "totalsales": "sales",
    "ยอดสุทธิ": "net_sales", "ยอดขายสุทธิ": "net_sales", "netsales": "net_sales", "net": "net_sales",
    "ออเดอ": "orders", "ออเดอร์": "orders", "order": "orders", "orders": "orders",
    "จำนวนบิล": "orders", "จำนวนออเดอร์": "orders", "บิล": "orders", "bills": "orders",
    "จำนวน": "qty", "จำนวนชิ้น": "qty", "qty": "qty", "quantity": "qty",
    "ส่วนลด": "discount", "discount": "discount",
    "%": "pct_source", "สัดส่วน": "pct_source", "percent": "pct_source",
}
STD_COLUMNS = set(COLUMN_ALIASES.values())
TEXT_COLS = {"date", "weekday", "category", "slot", "product"}

# หัวตาราง (คอลัมน์แรก) ที่ใช้ "หาตาราง" -> ชนิดตาราง  (คีย์ต้องเป็น casefold)
TABLE_ANCHORS = {
    "วันที่": "daily", "date": "daily",
    "หมวดหมู่": "category", "หมวด": "category", "หมวดสินค้า": "category",
    "ประเภทสินค้า": "category", "category": "category",
    "ช่วงเวลา": "slot", "เวลา": "slot", "time slot": "slot", "timeslot": "slot",
    "สินค้า": "product", "รายการสินค้า": "product", "product": "product",
}
KIND_KEY_COLUMN = {"daily": "date", "category": "category", "slot": "slot", "product": "product"}
TOTAL_MARKERS = {"รวม", "รวมทั้งหมด", "รวมทั้งสิ้น", "total", "sum", "grand total"}

# ตัวเลขสรุปบนหัวรายงาน (ข้อความในเซลล์แบบ "ป้ายชื่อ\nตัวเลข")
KPI_ALIASES = [
    ("ยอดขายรวม", "total_sales"), ("ส่วนลด", "discount"), ("เฉลี่ย", "avg_bill"),
    ("รายการขาย", "line_items"), ("ออเดอ", "orders_header"),
]

# คอลัมน์ที่ "ควรมี" ในแต่ละตาราง (ใช้แจ้งเตือนเมื่อเดือนใดขาด field)
EXPECTED_COLS = {
    "daily": {"sales", "orders"}, "category": {"sales", "qty"},
    "slot": {"sales", "orders"}, "product": {"sales", "qty"},
}
# ตารางที่แต่ละโหมดต้องใช้
MODE_TABLES = {"quarter": {"category", "slot"}, "weekday": {"daily"}, "week": {"daily"}, "day": {"daily"}}
MODE_PREFIX = {"quarter": "Quarter", "weekday": "Weekday", "week": "Week", "day": "Daily"}
ALL_MODES = ["quarter", "weekday", "week", "day"]

TOTAL_LABEL_Q = "รวมทั้งไตรมาส"
TOTAL_LABEL_ALL = "รวมทั้งหมด"
TOTAL_LABELS = {TOTAL_LABEL_Q, TOTAL_LABEL_ALL}
C_QTOTAL, C_QSHARE = "ยอดขายรวมไตรมาส", "สัดส่วนไตรมาส"
C_STOTAL, C_SSHARE, C_SAVG = "ยอดขายรวม", "สัดส่วนยอดขาย", "ยอดเฉลี่ย/วันที่เปิดขาย"
RENAME_TOTALS = {"qty": "จำนวนรวม", "orders": "ออเดอร์รวม", "net_sales": "ยอดสุทธิรวม",
                 "discount": "ส่วนลดรวม"}
RENAME_DAY = {"qty": "จำนวน", "orders": "ออเดอร์", "net_sales": "ยอดสุทธิ", "discount": "ส่วนลด"}


# ----------------------------------------------------------------------------
# 2) โครงสร้างข้อมูล
# ----------------------------------------------------------------------------
@dataclass
class MonthData:
    path: Path
    sheet: str
    month: int | None = None
    year_be: int | None = None
    kpi: dict = field(default_factory=dict)
    tables: dict = field(default_factory=dict)       # kind -> DataFrame
    notes: list = field(default_factory=list)        # (ระดับ, ข้อความ)

    @property
    def year_ce(self):
        return None if self.year_be is None else self.year_be - 543

    @property
    def ym(self):
        return (self.year_ce, self.month)

    @property
    def tag(self):
        return f"{MONTH_ABBR_TH.get(self.month, '?')} {self.year_be or '?'} ({self.path.name})"

    def note(self, level: str, msg: str):
        add_note(self.notes, level, f"[{self.tag}] {msg}")


def add_note(notes: list, level: str, msg: str):
    notes.append((level, msg))
    getattr(log, "warning" if level == "WARN" else "info")(msg)


# ----------------------------------------------------------------------------
# 3) ตัวช่วยเล็ก ๆ
# ----------------------------------------------------------------------------
def norm_text(s) -> str:
    """NFC + ตัดช่องว่างซ้ำ"""
    return re.sub(r"\s+", " ", unicodedata.normalize("NFC", str(s))).strip()


def norm_key(s) -> str:
    """คีย์สำหรับจับคู่ชื่อ (ไม่สนตัวพิมพ์เล็ก/ใหญ่และช่องว่างซ้ำ) เช่น 'wine' = 'Wine'"""
    return norm_text(s).casefold()


def anchor_kind(v) -> str | None:
    return TABLE_ANCHORS.get(norm_key(v)) if isinstance(v, str) else None


def to_number(x) -> float:
    if isinstance(x, bool) or x is None:
        return float("nan")
    if isinstance(x, (int, float)):
        return float(x)
    s = re.sub(r"[,\s%฿บาท]", "", str(x))
    try:
        return float(s)
    except ValueError:
        return float("nan")


_ARITH = re.compile(r"^[\d\s+\-*/().,]+$")


def safe_eval(formula: str):
    """คำนวณสูตรเลขคณิตง่าย ๆ เช่น '=8095+1419' เมื่อไฟล์ไม่มีค่า cache (ไม่รองรับ SUM ฯลฯ)"""
    expr = formula.lstrip("=").replace(",", "")
    if "**" in expr or not _ARITH.match(expr):
        return None
    try:
        return eval(expr, {"__builtins__": {}}, {})  # noqa: S307 (whitelist เฉพาะตัวเลข/เครื่องหมาย)
    except Exception:
        return None


def month_from_token(token: str) -> int | None:
    t = norm_text(token).casefold().strip(".")
    if not t:
        return None
    for m, names in THAI_MONTHS.items():
        if t.startswith(names[0]) or t == names[1].casefold().strip("."):
            return m                      # รองรับพิมพ์ผิดแบบ "กรกฎาคมคม"
    for m, names in EN_MONTHS.items():
        if t == names[-1] or t == names[0] or t.startswith(names[0]) or t in names:
            return m
    close = difflib.get_close_matches(t, [n[0] for n in THAI_MONTHS.values()], n=1, cutoff=0.75)
    if close:
        return next(m for m, n in THAI_MONTHS.items() if n[0] == close[0])
    return None


def slot_label(s) -> tuple[str, int | None]:
    """'19-20' / '19:00-20:00' / '23-00' -> ('19-20', ชั่วโมงเริ่ม)  ('23-00' และ '23-24' เป็นช่วงเดียวกัน)"""
    m = re.match(r"\s*(\d{1,2})(?:[:.]\d{2})?\s*(?:-|–|—|ถึง)\s*(\d{1,2})", str(s))
    if not m:
        return norm_text(s), None
    h1, h2 = int(m.group(1)) % 24, int(m.group(2))
    if h2 == 0:
        h2 = 24
    return f"{h1:02d}-{h2:02d}", h1


def infer_start_hour(hours: list[int]) -> int:
    """หาชั่วโมงเริ่มวันทำการเอง = ชั่วโมงหลัง 'ช่องว่างที่ยาวที่สุด' ของชั่วโมงที่พบ (18,19..23,0,1 -> 18)"""
    hs = sorted(set(hours))
    if len(hs) < 2:
        return hs[0] if hs else 0
    best_gap, start = -1, hs[0]
    for i, h in enumerate(hs):
        nxt = hs[(i + 1) % len(hs)]
        gap = (nxt - h) % 24 or 24
        if gap > best_gap:
            best_gap, start = gap, nxt
    return start


def quarter_of(month: int, year_ce: int, fy_start: int) -> tuple[int, int]:
    """คืน (ปีของไตรมาส (ค.ศ.), เลขไตรมาส) รองรับปีงบประมาณ"""
    q = ((month - fy_start) % 12) // 3 + 1
    fy = year_ce + 1 if (fy_start > 1 and month >= fy_start) else year_ce
    return fy, q


def be_str(ts) -> str:
    return "" if pd.isna(ts) else f"{ts.day:02d}/{ts.month:02d}/{ts.year + 543}"


def safe_concat(frames: list[pd.DataFrame]) -> pd.DataFrame:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", FutureWarning)
        return pd.concat(frames, ignore_index=True)


def value_columns(df: pd.DataFrame, exclude: Iterable[str] = ()) -> list[str]:
    """คอลัมน์ตัวเลขที่ต้องรวม (ไม่รวมคอลัมน์ข้อความ/ภายใน)"""
    skip = TEXT_COLS | {"pct_source"} | set(exclude)
    return [c for c in df.columns
            if c not in skip and not str(c).startswith("_") and pd.api.types.is_numeric_dtype(df[c])]


def total_name(c: str, table: dict = RENAME_TOTALS) -> str:
    return table.get(c) or (f"{c[2:]} (รวม)" if c.startswith("x_") else c)


def day_name(c: str) -> str:
    return RENAME_DAY.get(c) or (c[2:] if c.startswith("x_") else c)


def apply_aliases(path: str | Path):
    """โหลด alias เพิ่มเติมจาก JSON (ดูตัวอย่างในหัวไฟล์)"""
    try:
        data = json.loads(Path(path).read_text(encoding="utf-8"))
    except (OSError, ValueError) as e:
        raise SystemExit(f"อ่านไฟล์ alias '{path}' ไม่ได้: {e}")
    for name, std in data.get("columns", {}).items():
        if std not in STD_COLUMNS:
            raise SystemExit(f"alias คอลัมน์ '{name}' -> '{std}' ไม่ถูกต้อง (ใช้ได้: {sorted(STD_COLUMNS)})")
        COLUMN_ALIASES[norm_key(name).replace(" ", "")] = std
    for name, kind in data.get("tables", {}).items():
        if kind not in KIND_KEY_COLUMN:
            raise SystemExit(f"alias ตาราง '{name}' -> '{kind}' ไม่ถูกต้อง (ใช้ได้: {sorted(KIND_KEY_COLUMN)})")
        TABLE_ANCHORS[norm_key(name)] = kind
        COLUMN_ALIASES[norm_key(name).replace(" ", "")] = KIND_KEY_COLUMN[kind]


# ----------------------------------------------------------------------------
# 4) อ่านไฟล์ -> grid {(row, col): value}
# ----------------------------------------------------------------------------
def load_grids(path: Path):
    """คืน (ชื่อชีต, grid, สูตรที่คำนวณค่าไม่ได้) ของทุกชีต"""
    wb_f = load_workbook(path)                    # สูตร
    wb_v = load_workbook(path, data_only=True)    # ค่า cache (ว่างถ้าไฟล์ไม่เคยถูก Excel คำนวณ)
    try:
        for ws_f in wb_f.worksheets:
            ws_v = wb_v[ws_f.title]
            grid, uncached = {}, []
            for row in ws_f.iter_rows():
                for cell in row:
                    raw = cell.value
                    if raw is None:
                        continue
                    val = ws_v[cell.coordinate].value
                    if val is None:
                        if isinstance(raw, str) and raw.startswith("="):
                            val = safe_eval(raw)
                            if val is None:
                                uncached.append((cell.row, cell.column))
                                continue
                        else:
                            val = raw
                    if isinstance(val, str) and not val.strip():
                        continue
                    grid[(cell.row, cell.column)] = val
            yield ws_f.title, grid, uncached
    finally:
        wb_f.close()
        wb_v.close()


# ----------------------------------------------------------------------------
# 5) แกะเดือน/ปี, KPI, และตารางต่าง ๆ
# ----------------------------------------------------------------------------
def detect_period(md: MonthData, grid: dict):
    # 5.1 จากหัวรายงาน
    for (r, c), v in sorted(grid.items()):
        if r > 6 or not isinstance(v, str):
            continue
        m = re.search(r"ประจำเดือน\s*([^\d\s(]+)\s*(\d{4})?", v)
        if m:
            md.month = month_from_token(m.group(1))
            if m.group(2):
                md.year_be = int(m.group(2))
            break
    # 5.2 จากชื่อไฟล์/ชื่อชีต
    if md.month is None or md.year_be is None:
        for token in re.split(r"[^0-9A-Za-zก-๙.]+", f"{md.path.stem} {md.sheet}"):
            if md.month is None:
                md.month = month_from_token(token)
            if md.year_be is None and re.fullmatch(r"(25|20)\d{2}", token):
                y = int(token)
                md.year_be = y if y >= 2400 else y + 543


def parse_kpis(md: MonthData, grid: dict):
    for (r, c), v in grid.items():
        if r > 5 or not isinstance(v, str):
            continue
        lines = [ln for ln in v.splitlines() if ln.strip()]
        if len(lines) < 2:
            continue
        label = re.sub(r"[^ก-๙A-Za-z/ ]", "", lines[0]).strip()
        num = re.search(r"\d[\d,]*(?:\.\d+)?", " ".join(lines[1:]))
        if not label or not num:
            continue
        name = next((std for kw, std in KPI_ALIASES if kw in label), f"kpi_{label}")
        md.kpi[name] = to_number(num.group(0))


def read_block(grid: dict, hr: int, hc: int, max_row: int) -> tuple[pd.DataFrame, list[str]]:
    """อ่านตารางที่หัวคอลัมน์เริ่มที่ (hr, hc) ลงไปจนเจอแถว 'รวม' / แถวว่างติดกัน
    คืน (DataFrame, ชื่อคอลัมน์ข้อความที่ไม่รู้จักและถูกข้าม)"""
    cols = []
    c = hc
    while (hr, c) in grid:
        cols.append((c, str(grid[(hr, c)])))
        c += 1
    rows, blanks = [], 0
    for r in range(hr + 1, max_row + 1):
        key = grid.get((r, hc))
        if key is None:
            blanks += 1
            if blanks >= 2:
                break
            continue
        blanks = 0
        k = norm_text(key)
        if k.casefold() in TOTAL_MARKERS or anchor_kind(k):
            break
        rows.append({cc: grid.get((r, cc)) for cc, _ in cols})
    df = pd.DataFrame(rows, columns=[cc for cc, _ in cols])      # คีย์ด้วยเลขคอลัมน์ ชื่อซ้ำจึงไม่ทับกัน

    new, seen = [], set()
    for _, name in cols:
        std = COLUMN_ALIASES.get(norm_key(name).replace(" ", ""), "x_" + norm_text(name))
        while std in seen:
            std += "_dup"
        seen.add(std)
        new.append(std)
    df.columns = new

    dropped = []
    for col in list(df.columns):
        if col in TEXT_COLS:
            df[col] = df[col].map(
                lambda x: norm_text(x) if x is not None and not isinstance(x, (date, datetime)) else x)
            continue
        nums = df[col].map(to_number)
        nonnull = int(df[col].notna().sum())
        if col.startswith("x_") and nonnull and nums.notna().sum() < 0.5 * nonnull:
            dropped.append(col[2:])                     # คอลัมน์ข้อความที่ไม่รู้จัก (เช่น หมายเหตุ)
            df = df.drop(columns=col)
        else:
            df[col] = nums
    return df, dropped


def parse_tables(md: MonthData, grid: dict, required: set[str]):
    max_row = max(r for r, _ in grid)
    anchors = {}
    for (r, c), v in sorted(grid.items()):
        kind = anchor_kind(v)
        if kind is None:
            continue
        if kind in anchors:
            if kind in required:
                md.note("WARN", f"พบตาราง '{kind}' มากกว่า 1 ที่ ใช้ตารางแรก (แถว {anchors[kind][0]})")
        else:
            anchors[kind] = (r, c)
    for kind, (hr, hc) in anchors.items():
        df, dropped = read_block(grid, hr, hc, max_row)
        if df.empty:
            if kind in required:
                md.note("WARN", f"ตาราง '{kind}' ไม่มีข้อมูล")
            continue
        if "sales" not in df.columns and "net_sales" in df.columns:
            df["sales"] = df["net_sales"]
            if kind in required:
                md.note("WARN", f"ตาราง '{kind}' ไม่มีคอลัมน์ยอดขาย ใช้ 'ยอดสุทธิ' แทน")
        md.tables[kind] = df
        if kind not in required:
            continue
        missing = EXPECTED_COLS[kind] - set(df.columns)
        extra = [c for c in df.columns if c not in EXPECTED_COLS[kind] | TEXT_COLS | {"pct_source"}]
        if missing:
            md.note("WARN", f"ตาราง '{kind}' ขาด field: {sorted(missing)} (นับเป็น 0)")
        if extra:
            md.note("INFO", f"ตาราง '{kind}' มี field เพิ่ม: {[c[2:] if c.startswith('x_') else c for c in extra]} "
                            f"(รวมให้ด้วย)")
        if dropped:
            md.note("INFO", f"ตาราง '{kind}' ข้ามคอลัมน์ข้อความที่ไม่รู้จัก: {dropped}")
        bad = int(df.drop(columns=[c for c in df.columns if c in TEXT_COLS]).isna().sum().sum())
        if bad:
            md.note("WARN", f"ตาราง '{kind}' มีเซลล์ตัวเลขว่าง/อ่านไม่ได้ {bad} เซลล์ (นับเป็น 0)")
    for kind in sorted(required):
        if kind not in md.tables:
            md.note("WARN", f"ไม่พบตาราง '{kind}' ในชีตนี้ (ยอดส่วนนี้นับเป็น 0)")


def finalize_daily(md: MonthData):
    """แปลง 'dd/mm' เป็นวันที่จริง, ตรวจเดือนตรงหัวรายงาน, และคำนวณ 'วัน' (จ.-อา.) จากวันที่จริง"""
    df = md.tables.get("daily")
    if df is None:
        return
    parsed = []
    for v in (df["date"] if "date" in df else [None] * len(df)):
        if isinstance(v, (date, datetime)):
            parsed.append((v.day, v.month, v.year))
            continue
        m = re.match(r"\s*(\d{1,2})/(\d{1,2})(?:/(\d{2,4}))?", str(v))
        if not m:
            parsed.append((None, None, None))
            continue
        y = int(m.group(3)) if m.group(3) else None
        if y is not None and y >= 2400:
            y -= 543
        parsed.append((int(m.group(1)), int(m.group(2)), y))
    months = Counter(p[1] for p in parsed if p[1])
    if months:
        dominant = months.most_common(1)[0][0]
        if md.month is None:
            md.month = dominant
        elif md.month != dominant:
            md.note("WARN", f"เดือนในหัวรายงาน ({md.month}) ไม่ตรงกับวันที่ในตาราง ({dominant}) ใช้ค่าจากหัวรายงาน")
    dates = []
    for d, m, y in parsed:
        try:
            dates.append(pd.Timestamp(date(y or md.year_ce or datetime.now().year, m, d)))
        except (TypeError, ValueError):
            dates.append(pd.NaT)
    df["date"] = pd.to_datetime(pd.Series(dates, index=df.index))

    given = df["weekday"] if "weekday" in df else pd.Series([None] * len(df), index=df.index)
    names, mism = [], 0
    for d, g in zip(df["date"], given):
        gi = WEEKDAY_ALIASES.get(norm_key(g).strip(".")) if isinstance(g, str) else None
        if pd.isna(d):
            idx = gi
        else:
            idx = d.dayofweek
            if gi is not None and gi != idx:
                mism += 1
        names.append(THAI_WEEKDAYS[idx] if idx is not None else None)
    df["weekday"] = names
    if mism:
        md.note("WARN", f"คอลัมน์ 'วัน' ไม่ตรงกับวันที่จริง {mism} แถว ใช้วันที่คำนวณจากวันที่")
    md.tables["daily"] = df


def parse_workbook(path: Path, required: set[str]) -> list[MonthData]:
    out = []
    for sheet, grid, uncached in load_grids(path):
        if not grid:
            continue
        md = MonthData(path=path, sheet=sheet)
        detect_period(md, grid)
        parse_kpis(md, grid)
        parse_tables(md, grid, required)
        if not md.tables:
            log.info("ข้ามชีต '%s' ใน %s (ไม่พบตารางที่รู้จัก)", sheet, path.name)
            continue
        total_rows = {r for (r, _), v in grid.items() if isinstance(v, str) and norm_key(v) in TOTAL_MARKERS}
        lost = [(r, c) for r, c in uncached if r not in total_rows]
        if lost:
            r, c = lost[0]
            md.note("WARN", f"มีสูตรที่ไม่มีค่าคำนวณ {len(lost)} เซลล์ (เช่น {get_column_letter(c)}{r}) ถูกข้าม "
                            f"-> เปิดไฟล์ใน Excel แล้วกด Save ก่อนรัน")
        finalize_daily(md)
        if md.month is None:
            md.note("WARN", "ระบุเดือนไม่ได้ ข้ามชีตนี้")
            continue
        if md.year_be is None:
            md.year_be = datetime.now().year + 543
            md.note("WARN", f"ระบุปีไม่ได้ ใช้ปีปัจจุบัน {md.year_be}")
        out.append(md)
    return out


# ----------------------------------------------------------------------------
# 6) รวมข้อมูล -- โหมด quarter (หมวดหมู่ / ช่วงเวลา)
# ----------------------------------------------------------------------------
def month_labels(months: list[MonthData]) -> dict:
    multi = len({md.year_be for md in months}) > 1
    return {md.ym: (f"{MONTH_ABBR_TH[md.month]} {md.year_be}" if multi else MONTH_ABBR_TH[md.month])
            for md in months}


def period_label(months: list[MonthData]) -> str:
    first, last = months[0], months[-1]
    if first.ym == last.ym:
        return f"{MONTH_ABBR_TH[first.month]} {first.year_be}  (1 เดือน)"
    start = f"{MONTH_ABBR_TH[first.month]}" + (f" {first.year_be}" if first.year_be != last.year_be else "")
    return f"{start} – {MONTH_ABBR_TH[last.month]} {last.year_be}  ({len(months)} เดือน)"


def merge_dimension(months: list[MonthData], kind: str, key: str, labels: dict, slot_start=None):
    """รวมตาราง `kind` ทุกเดือน -> ยอดต่อเดือน, ยอดรวม, สัดส่วน (%)  (ยังไม่มีแถวรวม)"""
    parts = []
    for md in months:
        df = md.tables.get(kind)
        if df is None or key not in df or "sales" not in df:
            continue
        d = df.copy()
        d["_m"] = labels[md.ym]
        parts.append(d)
    if not parts:
        return None
    allr = safe_concat(parts)                 # concat = union ของ field (field ที่ขาดจะเป็น NaN)
    if kind == "slot":
        lab = allr[key].map(slot_label)
        hours = [t[1] for t in lab if t[1] is not None]
        start = slot_start if isinstance(slot_start, int) else infer_start_hour(hours)
        allr["_order"] = [(h - start) % 24 if h is not None else 99 for _, h in lab]
        allr[key] = [t[0] for t in lab]
    allr["_k"] = allr[key].map(norm_key)
    first_seen = allr.drop_duplicates("_k").set_index("_k")      # ใช้การสะกดที่เจอครั้งแรก

    month_cols = [labels[md.ym] for md in months]
    num_cols = value_columns(allr)
    totals = allr.groupby("_k", sort=False)[num_cols].sum(min_count=1).fillna(0)
    by_month = (allr.pivot_table(index="_k", columns="_m", values="sales", aggfunc="sum", fill_value=0)
                .reindex(columns=month_cols, fill_value=0).reindex(totals.index).fillna(0))

    res = pd.DataFrame({key: first_seen[key].reindex(totals.index).values}, index=totals.index)
    for m in month_cols:
        res[m] = by_month[m]
    res[C_QTOTAL] = totals["sales"]
    grand = res[C_QTOTAL].sum()
    res[C_QSHARE] = res[C_QTOTAL] / grand if grand else 0.0
    for c in num_cols:
        if c != "sales":
            res[total_name(c)] = totals[c]

    if kind == "slot":
        res = res.assign(_o=first_seen["_order"].reindex(res.index).values).sort_values("_o").drop(columns="_o")
    else:
        res = res.sort_values(C_QTOTAL, ascending=False)
    return res.reset_index(drop=True)


def add_total_row(df: pd.DataFrame, total_label: str = TOTAL_LABEL_Q) -> pd.DataFrame:
    total = {df.columns[0]: total_label}
    for c in df.columns[1:]:
        total[c] = (1.0 if df[c].sum() else 0.0) if str(c).startswith("สัดส่วน") else df[c].sum()
    return pd.concat([df, pd.DataFrame([total])], ignore_index=True)


def build_monthly_check(months: list[MonthData], tol: float, notes: list) -> pd.DataFrame:
    """ยอดรวมของแต่ละเดือนจากแหล่งต่าง ๆ + สถานะการตรวจยอด"""
    rows = []
    for md in months:
        def tsum(kind):
            df = md.tables.get(kind)
            return float(df["sales"].sum()) if df is not None and "sales" in df else float("nan")
        vals = {"ยอดบนหัวรายงาน": md.kpi.get("total_sales", float("nan")), "ยอดรายวัน": tsum("daily"),
                "ยอดหมวดหมู่": tsum("category"), "ยอดช่วงเวลา": tsum("slot")}
        have = {k: v for k, v in vals.items() if not pd.isna(v)}
        spread = max(have.values()) - min(have.values()) if len(have) >= 2 else float("nan")
        if len(have) < 2:
            status = "ข้อมูลไม่พอเทียบ"
        elif spread <= tol:
            status = "ตรงกัน"
        else:
            status = "ไม่ตรง (ตรวจไฟล์ต้นทาง)"
            detail = ", ".join(f"{k}={v:,.0f}" for k, v in have.items())
            add_note(notes, "WARN", f"[{MONTH_ABBR_TH[md.month]} {md.year_be}] ยอดจากแต่ละส่วนไม่ตรงกัน: {detail} "
                                    f"(สัดส่วนคำนวณจากยอดรวมของตารางนั้น ๆ)")
        rows.append({"เดือน": f"{MONTH_FULL_TH[md.month]} {md.year_be}", **vals,
                     "ผลต่างสูงสุด": spread, "สถานะ": status})
        daily = md.tables.get("daily")
        if daily is not None and "orders" in daily:
            for k in ("orders_header", "line_items"):
                if k in md.kpi and abs(md.kpi[k] - daily["orders"].sum()) > tol:
                    md.note("INFO", f"ตัวเลข '{k}' บนหัวรายงาน = {md.kpi[k]:,.0f} "
                                    f"ไม่เท่ากับผลรวมออเดอร์รายวัน = {daily['orders'].sum():,.0f} (นิยามอาจต่างกัน)")
    return pd.DataFrame(rows)


def kpi_quarter(period: str, cat_body, slot_body, monthly: pd.DataFrame) -> pd.DataFrame:
    items = [("ช่วงข้อมูล", period)]
    if cat_body is not None:
        items.append(("ยอดขายรวมตามหมวดหมู่ (บาท)", cat_body[C_QTOTAL].sum()))
    if slot_body is not None:
        items.append(("ยอดขายรวมตามช่วงเวลา (บาท)", slot_body[C_QTOTAL].sum()))
    if monthly["ยอดบนหัวรายงาน"].notna().any():
        items.append(("ยอดขายรวมตามหัวรายงาน (บาท)", monthly["ยอดบนหัวรายงาน"].sum()))
    if cat_body is not None and len(cat_body):
        c = cat_body.iloc[0]
        items.append(("หมวดหมู่อันดับ 1", f"{c.iloc[0]}  ({c[C_QSHARE]:.1%} / {c[C_QTOTAL]:,.0f} บาท)"))
    if slot_body is not None and len(slot_body):
        s = slot_body.loc[slot_body[C_QTOTAL].idxmax()]
        items.append(("ช่วงเวลาที่ขายดีที่สุด", f"{s.iloc[0]}  ({s[C_QSHARE]:.1%} / {s[C_QTOTAL]:,.0f} บาท)"))
    return pd.DataFrame(items, columns=["รายการ", "ค่า"])


# ----------------------------------------------------------------------------
# 7) รวมข้อมูล -- โหมด weekday / week / day (จากตารางรายวัน)
# ----------------------------------------------------------------------------
def collect_daily(months: list[MonthData], labels: dict, notes: list, need_date: bool) -> pd.DataFrame | None:
    parts = []
    for md in months:
        df = md.tables.get("daily")
        if df is None or "sales" not in df:
            continue
        d = df.copy()
        d["_mlabel"] = labels[md.ym]
        parts.append(d)
    if not parts:
        return None
    d = safe_concat(parts)
    col = "date" if need_date else "weekday"
    bad = int(d[col].isna().sum())
    if bad:
        add_note(notes, "WARN", f"มี {bad} แถวในตารางรายวันที่ระบุ{'วันที่' if need_date else 'วัน'}ไม่ได้ "
                                f"ถูกข้าม (ยอดรวมจึงอาจต่ำกว่าจริง)")
        d = d[d[col].notna()]
    return d.reset_index(drop=True)


def aggregate_daily(d: pd.DataFrame, gcol: str, order: list, key_name: str, label_fn, info_fn=None) -> pd.DataFrame:
    extra = value_columns(d, exclude={"sales"})
    rows, info_cols = [], []
    for g in order:
        sub = d[d[gcol] == g]
        row = {key_name: label_fn(g)}
        if info_fn:
            row.update(info_fn(g, sub))
            info_cols = [c for c in row if c != key_name]
        row["จำนวนวันที่มีข้อมูล"] = len(sub)
        row["วันที่เปิดขาย"] = int((sub["sales"] > 0).sum())
        row[C_STOTAL] = sub["sales"].sum()
        for c in extra:
            row[total_name(c)] = sub[c].sum()
        rows.append(row)
    df = pd.DataFrame(rows)
    grand = df[C_STOTAL].sum()
    pos = df.columns.get_loc(C_STOTAL) + 1
    df.insert(pos, C_SSHARE, df[C_STOTAL] / grand if grand else 0.0)
    df.insert(pos + 1, C_SAVG, [s / n if n else 0.0 for s, n in zip(df[C_STOTAL], df["วันที่เปิดขาย"])])

    total = {c: "" for c in df.columns}
    total[key_name] = TOTAL_LABEL_ALL
    for c in df.columns:
        if c in (key_name, C_SSHARE, C_SAVG) or c in info_cols:
            continue
        total[c] = df[c].sum()
    total[C_SSHARE] = 1.0 if grand else 0.0
    total[C_SAVG] = grand / total["วันที่เปิดขาย"] if total["วันที่เปิดขาย"] else 0.0
    return pd.concat([df, pd.DataFrame([total])], ignore_index=True)


def build_weekday_table(d: pd.DataFrame) -> pd.DataFrame:
    d = d.assign(_wd=d["weekday"].map({n: i for i, n in enumerate(THAI_WEEKDAYS)}))
    return aggregate_daily(d, "_wd", list(range(7)), "วัน", lambda i: THAI_WEEKDAYS[i])


def build_week_table(d: pd.DataFrame) -> pd.DataFrame:
    d = d.assign(_wk=d["date"] - pd.to_timedelta(d["date"].dt.dayofweek, unit="D"))
    order = sorted(d["_wk"].unique())

    def label(g):
        g = pd.Timestamp(g).to_pydatetime()
        return f"{be_str(g)} – {be_str(g + timedelta(days=6))}"

    def info(g, sub):
        n = len(sub)
        return {"สัปดาห์ที่ (ISO)": int(pd.Timestamp(g).isocalendar()[1]),
                "หมายเหตุ": "" if n >= 7 else f"มีข้อมูลเพียง {n} วัน (สัปดาห์คาบเกี่ยวช่วงข้อมูล)"}

    return aggregate_daily(d, "_wk", order, "สัปดาห์ (จ.–อา.)", label, info)


def build_day_table(d: pd.DataFrame) -> pd.DataFrame:
    d = d.sort_values("date").reset_index(drop=True)
    extra = value_columns(d, exclude={"sales"})
    out = pd.DataFrame({"เดือน": d["_mlabel"], "วันที่ (พ.ศ.)": d["date"].map(be_str),
                        "วันที่ (ค.ศ.)": d["date"].dt.date, "วัน": d["weekday"], "ยอดขาย": d["sales"]})
    grand = out["ยอดขาย"].sum()
    out[C_SSHARE] = out["ยอดขาย"] / grand if grand else 0.0
    for c in extra:
        out[day_name(c)] = d[c]
    total = {"เดือน": TOTAL_LABEL_ALL, "ยอดขาย": grand, C_SSHARE: 1.0 if grand else 0.0}
    for c in extra:
        total[day_name(c)] = d[c].sum()
    return pd.concat([out, pd.DataFrame([total])], ignore_index=True)


def kpi_daily(mode: str, d: pd.DataFrame, table: pd.DataFrame, period: str) -> pd.DataFrame:
    body = table.iloc[:-1]
    total = d["sales"].sum()
    open_days = int((d["sales"] > 0).sum())
    items = [("ช่วงข้อมูล", period), ("ยอดขายรวมทั้งหมด (บาท)", total)]
    if "orders" in d.columns:
        items.append(("ออเดอร์รวม (ตามตารางรายวัน)", d["orders"].sum()))
    items += [("จำนวนวันที่มีข้อมูล", len(d)), ("จำนวนวันที่เปิดขาย (ยอด > 0)", open_days),
              ("ยอดเฉลี่ยต่อวันที่เปิดขาย (บาท)", total / open_days if open_days else float("nan"))]
    if mode == "day":
        if body["ยอดขาย"].notna().any():
            r = body.loc[body["ยอดขาย"].idxmax()]
            items.append(("วันที่ยอดสูงสุด", f"{r['วันที่ (พ.ศ.)']} {r['วัน']}  "
                                             f"({r['ยอดขาย']:,.0f} บาท / {r[C_SSHARE]:.1%})"))
    else:
        what = "วันในสัปดาห์" if mode == "weekday" else "สัปดาห์"
        if body[C_STOTAL].notna().any():
            r = body.loc[body[C_STOTAL].idxmax()]
            items.append((f"{what}ที่ยอดรวมสูงสุด", f"{r.iloc[0]}  ({r[C_STOTAL]:,.0f} บาท / {r[C_SSHARE]:.1%})"))
        if mode == "weekday" and (body["วันที่เปิดขาย"] > 0).any():
            ok = body[body["วันที่เปิดขาย"] > 0]
            r = ok.loc[ok[C_SAVG].idxmax()]
            items.append(("วันที่ยอดเฉลี่ยต่อวันสูงสุด", f"{r.iloc[0]}  ({r[C_SAVG]:,.0f} บาท/วัน)"))
    return pd.DataFrame(items, columns=["รายการ", "ค่า"])


# ----------------------------------------------------------------------------
# 8) เขียน Excel / CSV
# ----------------------------------------------------------------------------
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
TOTAL_FILL = PatternFill("solid", fgColor="D9E1F2")
BASE_FONT = "Tahoma"


def _is_num(v) -> bool:
    return pd.api.types.is_number(v) and not isinstance(v, bool) and not pd.isna(v)   # รวม numpy int/float


def write_table(writer, sheet: str, df: pd.DataFrame, startrow: int = 0, title: str | None = None) -> int:
    """เขียน DataFrame + จัดรูปแบบ คืนแถวถัดไปที่ว่าง"""
    df = df.drop(columns=[c for c in df.columns if str(c).startswith("_")])
    if title:
        startrow += 1
    df.to_excel(writer, sheet_name=sheet, index=False, startrow=startrow)
    ws = writer.sheets[sheet]
    if title:
        ws.cell(startrow, 1, title).font = Font(name=BASE_FONT, bold=True, size=12)
    for j, col in enumerate(df.columns, start=1):
        h = ws.cell(startrow + 1, j)
        h.font, h.fill = Font(name=BASE_FONT, bold=True, color="FFFFFF"), HEADER_FILL
        h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        pct = str(col).startswith("สัดส่วน")
        numeric_col = pd.api.types.is_numeric_dtype(df[col])
        decimals_col = numeric_col and bool((df[col].dropna() % 1 != 0).any())
        for i in range(len(df)):
            cell = ws.cell(startrow + 2 + i, j)
            is_total = str(df.iloc[i, 0]) in TOTAL_LABELS
            cell.font = Font(name=BASE_FONT, bold=is_total)
            if is_total:
                cell.fill = TOTAL_FILL
            v = df.iloc[i, j - 1]
            if pct and _is_num(v):
                cell.number_format = "0.0%"
            elif _is_num(v):
                decimals = decimals_col if numeric_col else float(v) % 1 != 0
                cell.number_format = "#,##0.00" if decimals else "#,##0"
        width = max([len(str(col)) * 1.2] + [len(str(v)) * 1.1 for v in df[col].head(60)]) + 2
        letter = get_column_letter(j)
        ws.column_dimensions[letter].width = max(ws.column_dimensions[letter].width or 0, min(max(width, 10), 60))
    return startrow + len(df) + 3


def write_workbook(path: Path, title: str, kpi: pd.DataFrame, monthly: pd.DataFrame, tables: list, notes: list):
    with pd.ExcelWriter(path, engine="openpyxl") as w:
        r = write_table(w, "สรุป", kpi, title=title)
        write_table(w, "สรุป", monthly, startrow=r, title="ยอดรายเดือนและการตรวจยอด")
        w.sheets["สรุป"].column_dimensions["A"].width = 44
        w.sheets["สรุป"].column_dimensions["B"].width = 44
        for name, df in tables:
            if df is not None:
                write_table(w, name, df)
                w.sheets[name].freeze_panes = "B2"
        seen, uniq = set(), []
        for n in notes:
            if n not in seen:
                seen.add(n)
                uniq.append(n)
        nt = pd.DataFrame(uniq, columns=["ระดับ", "ข้อสังเกต"]) if uniq else \
            pd.DataFrame([("INFO", "ไม่มีข้อสังเกต")], columns=["ระดับ", "ข้อสังเกต"])
        write_table(w, "ข้อสังเกต", nt)
        w.sheets["ข้อสังเกต"].column_dimensions["B"].width = 120


def export_csv(out_dir: Path, stem: str, kpi: pd.DataFrame, tables: list):
    kpi.to_csv(out_dir / f"{stem}_สรุป.csv", index=False, encoding="utf-8-sig")
    for name, df in tables:
        if df is not None:
            df.drop(columns=[c for c in df.columns if str(c).startswith("_")]).to_csv(
                out_dir / f"{stem}_{name}.csv", index=False, encoding="utf-8-sig")


def show_console(label: str, kpi: pd.DataFrame, tables: list, xlsx: Path):
    print(f"\n===== {label} =====")
    print(kpi.to_string(index=False))
    for name, df in tables:
        if df is None:
            continue
        keys = [c for c in df.columns if c in ("วันที่ (พ.ศ.)",)] or [df.columns[0]]
        show = df[keys + [c for c in df.columns
                          if c in (C_QTOTAL, C_QSHARE, C_STOTAL, C_SSHARE, C_SAVG, "ยอดขาย")]].copy()
        for c in show.columns[len(keys):]:
            show[c] = show[c].map(lambda v, c=c: "" if pd.isna(v) or v == "" else
                                  (f"{v:.1%}" if str(c).startswith("สัดส่วน") else f"{v:,.0f}"))
        print(f"\n-- {name} --\n" + (show.to_string(index=False) if len(show) <= 40 else
                                     show.head(40).to_string(index=False) + "\n..."))
    print(f"\n-> เขียนไฟล์: {xlsx}")


# ----------------------------------------------------------------------------
# 9) ประมวลผลแต่ละโหมด
# ----------------------------------------------------------------------------
def process_quarter(months, label, expected, out_dir: Path, tol, csv, slot_start) -> Path | None:
    months = sorted(months, key=lambda m: m.ym)
    notes = [n for md in months for n in md.notes]
    present = {md.month for md in months}
    missing = [MONTH_ABBR_TH[m] for m in expected if m not in present]
    if missing:
        add_note(notes, "WARN", f"ไตรมาสนี้ข้อมูลไม่ครบ ขาดเดือน: {', '.join(missing)} "
                                f"(ยอดรวมเป็นยอดเฉพาะเดือนที่มี)")
    labels = month_labels(months)
    cat = merge_dimension(months, "category", "category", labels)
    slot = merge_dimension(months, "slot", "slot", labels, slot_start)
    if cat is None and slot is None:
        log.error("ไตรมาส %s: ไม่พบทั้งตารางหมวดหมู่และช่วงเวลา ข้าม", label)
        return None
    monthly = build_monthly_check(months, tol, notes)
    period = period_label(months)
    kpi = kpi_quarter(period, cat, slot, monthly)
    tables = []
    for name, df, key in (("หมวดหมู่", cat, "category"), ("ช่วงเวลา", slot, "slot")):
        if df is not None:
            df = add_total_row(df).rename(columns={key: name})
        tables.append((name, df))
    out_dir.mkdir(parents=True, exist_ok=True)
    xlsx = out_dir / f"{MODE_PREFIX['quarter']}_{label.replace(' ', '_').replace('/', '-')}.xlsx"
    write_workbook(xlsx, f"สรุปยอดขาย {label}", kpi, monthly, tables, notes)
    if csv:
        export_csv(out_dir, xlsx.stem, kpi, tables)
    show_console(f"{label} | {period}", kpi, tables, xlsx)
    return xlsx


def process_daily_mode(mode: str, months, label: str, out_dir: Path, tol, csv) -> Path | None:
    months = sorted(months, key=lambda m: m.ym)
    notes = [n for md in months for n in md.notes]
    labels = month_labels(months)
    d = collect_daily(months, labels, notes, need_date=(mode != "weekday"))
    if d is None or d.empty:
        log.error("โหมด %s: ไม่พบตารางรายวันที่ใช้งานได้ ข้าม", mode)
        return None
    table = {"weekday": build_weekday_table, "week": build_week_table, "day": build_day_table}[mode](d)
    sheet = {"weekday": "ตามวันในสัปดาห์", "week": "รายสัปดาห์", "day": "รายวัน"}[mode]
    monthly = build_monthly_check(months, tol, notes)
    period = period_label(months)
    kpi = kpi_daily(mode, d, table, period)
    if mode != "day":
        add_note(notes, "INFO", "ตารางหมวดหมู่/ช่วงเวลาในไฟล์ต้นทางเป็นยอดรวมรายเดือน จึงแบ่งตามวัน/สัปดาห์ไม่ได้ "
                                "(ใช้โหมด quarter สำหรับสองตารางนี้)")
    tables = [(sheet, table)]
    out_dir.mkdir(parents=True, exist_ok=True)
    xlsx = out_dir / f"{MODE_PREFIX[mode]}_{label}.xlsx"
    write_workbook(xlsx, f"สรุปยอดขาย{sheet} {label}", kpi, monthly, tables, notes)
    if csv:
        export_csv(out_dir, xlsx.stem, kpi, tables)
    show_console(f"{sheet} | {period}", kpi, tables, xlsx)
    return xlsx


# ----------------------------------------------------------------------------
# 10) Orchestration
# ----------------------------------------------------------------------------
def collect_files(inputs: Iterable[str | Path]) -> list[Path]:
    skip = tuple(f"{p}_" for p in MODE_PREFIX.values()) + ("~$",)
    files = []
    for p in map(Path, inputs):
        if p.is_dir():
            files += sorted(f for f in p.glob("*.xlsx") if not f.name.startswith(skip))
        elif p.suffix.lower() in {".xlsx", ".xlsm"} and p.exists():
            files.append(p)
        else:
            log.warning("ข้าม %s (ไม่พบไฟล์/ไม่ใช่ .xlsx)", p)
    return files


def run_pipeline(inputs, output_dir="./output", modes=("quarter",), quarter: str | None = None,
                 fiscal_start_month: int = 1, tolerance: float = 1.0, csv: bool = False,
                 slot_start_hour: int | None = None, aliases: str | None = None) -> list[Path]:
    modes = ALL_MODES if "all" in modes else [m for m in ALL_MODES if m in modes]
    if not modes:
        raise SystemExit(f"ไม่รู้จักโหมด (ใช้ได้: {', '.join(ALL_MODES)}, all)")
    if aliases:
        apply_aliases(aliases)
    required = set().union(*(MODE_TABLES[m] for m in modes))

    files = collect_files(inputs)
    if not files:
        raise SystemExit("ไม่พบไฟล์ .xlsx ใน input")
    all_months: list[MonthData] = []
    for f in files:
        try:
            all_months += parse_workbook(f, required)
        except Exception as e:  # ไฟล์เสียไฟล์เดียวไม่ควรทำให้ทั้ง pipeline ล้ม
            log.error("อ่านไฟล์ %s ไม่สำเร็จ: %s", f.name, e)

    by_period: dict = {}                       # กันเดือนซ้ำ (เก็บไฟล์ที่แก้ไขล่าสุด)
    for md in all_months:
        k = md.ym
        if k in by_period:
            old = by_period[k]
            keep = md if md.path.stat().st_mtime >= old.path.stat().st_mtime else old
            log.warning("พบเดือนซ้ำ %s: %s และ %s -> ใช้ %s", MONTH_ABBR_TH[md.month], old.path.name,
                        md.path.name, keep.path.name)
            by_period[k] = keep
        else:
            by_period[k] = md

    groups: dict = {}
    for (y, m), md in by_period.items():
        groups.setdefault(quarter_of(m, y, fiscal_start_month), []).append(md)
    labelled = {f"{fy + 543}-Q{q}": (fy, q, mds) for (fy, q), mds in sorted(groups.items())}
    if quarter:
        labelled = {k: v for k, v in labelled.items() if k.upper() == quarter.upper()}
    if not labelled:
        found = ", ".join(f"{fy + 543}-Q{q}" for fy, q in sorted(groups)) or "-"
        raise SystemExit(f"ไม่พบข้อมูลของไตรมาส {quarter} (ไตรมาสที่พบ: {found})" if quarter
                         else "ไม่พบข้อมูลที่ใช้งานได้ -- ถ้าชื่อหัวตาราง/คอลัมน์ในไฟล์ไม่ตรงกับที่รู้จัก "
                              "ให้เพิ่มด้วย --aliases (ดูตัวอย่างในหัวไฟล์)")

    out_dir, outputs = Path(output_dir), []
    if "quarter" in modes:
        for label, (fy, q, mds) in labelled.items():
            start = (fiscal_start_month - 1 + (q - 1) * 3) % 12 + 1
            expected = [(start - 1 + i) % 12 + 1 for i in range(3)]
            outputs.append(process_quarter(mds, label, expected, out_dir, tolerance, csv, slot_start_hour))

    daily_modes = [m for m in modes if m != "quarter"]
    if daily_modes:
        selected = sorted((md for _, _, mds in labelled.values() for md in mds), key=lambda m: m.ym)
        if quarter:
            label = next(iter(labelled))
        else:
            a, b = selected[0], selected[-1]
            label = f"{a.year_be}-{a.month:02d}" + ("" if a.ym == b.ym else f"_{b.year_be}-{b.month:02d}")
        for mode in daily_modes:
            outputs.append(process_daily_mode(mode, selected, label, out_dir, tolerance, csv))

    outputs = [o for o in outputs if o]
    if not outputs:
        raise SystemExit("ไม่มีรายงานที่สร้างได้ (ตรวจข้อความ ERROR/WARNING ด้านบน)")
    return outputs


def main(argv=None):
    ap = argparse.ArgumentParser(description="สรุปรายงานยอดขายรายเดือนจาก POS: รายไตรมาส / วันในสัปดาห์ / รายสัปดาห์ / รายวัน")
    ap.add_argument("-i", "--input", nargs="+", default=["."], help="โฟลเดอร์ หรือไฟล์ .xlsx (ระบุได้หลายอัน)")
    ap.add_argument("-o", "--output-dir", default="./output")
    ap.add_argument("-m", "--mode", nargs="+", default=["quarter"], choices=ALL_MODES + ["all"],
                    help="โหมดสรุป เลือกได้หลายอัน: quarter(default) weekday week day all")
    ap.add_argument("-q", "--quarter", help="เลือกไตรมาส เช่น 2569-Q3 (ปี พ.ศ.) ไม่ระบุ = ทุกไตรมาสที่พบ")
    ap.add_argument("--fiscal-start-month", type=int, default=1, choices=range(1, 13), metavar="1-12",
                    help="เดือนเริ่มปีงบประมาณ (ค่าเริ่มต้น 1 = ไตรมาสปฏิทิน; ปีงบรัฐ = 10)")
    ap.add_argument("--slot-start-hour", type=int, default=None, choices=range(0, 24), metavar="0-23",
                    help="ชั่วโมงเริ่มวันทำการ ใช้เรียงช่วงเวลา (ไม่ระบุ = ตรวจจากข้อมูลเอง)")
    ap.add_argument("--tolerance", type=float, default=1.0, help="ผลต่างยอดที่ยอมรับได้ (บาท) ตอนตรวจยอด")
    ap.add_argument("--aliases", help="ไฟล์ JSON เพิ่ม alias ชื่อคอลัมน์/ชื่อตาราง (ดูตัวอย่างในหัวไฟล์)")
    ap.add_argument("--csv", action="store_true", help="เขียน CSV แยกแต่ละตารางด้วย")
    ap.add_argument("-v", "--verbose", action="store_true")
    a = ap.parse_args(argv)
    logging.basicConfig(level=logging.INFO if a.verbose else logging.WARNING, format="%(levelname)s %(message)s")
    run_pipeline(a.input, a.output_dir, a.mode, a.quarter, a.fiscal_start_month, a.tolerance, a.csv,
                 a.slot_start_hour, a.aliases)


if __name__ == "__main__":
    sys.exit(main())