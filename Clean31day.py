import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border
import re, shutil, sys, os

# ============================================================
# CONFIG — แก้ตรงนี้เพื่อใช้กับเดือนใหม่
# ============================================================
MONTH_LABEL   = 'สิงหาคม 2569'
MONTH_SHORT   = 'ส.ค.'
DATE_RANGE    = '01-31 ส.ค.'
TOTAL_DAYS    = 31
OUTPUT_FILE   = 'Data_August_2569.xlsx'

FILE_DAILY    = 'D31.xlsx' #รายวัน
FILE_CATEGORY = 'C31.xlsx' #หมวดหมู่
FILE_PRODUCTS = 'P31.xlsx' #สินค้า
FILE_HOURLY   = 'H31.xlsx' #รายชั่วโมง
FILE_RECEIPTS = 'R31.xlsx' #ตามบิล
TEMPLATE_FILE = './Template/Template31day.xlsx'

# ============================================================
# Helpers
# ============================================================
def parse_thb(val):
    if pd.isna(val): return 0.0
    return float(str(val).replace('฿','').replace(',','').strip() or 0)

def fmt_hour(s):
    m = re.match(r'(\d+):00-(\d+):00', str(s))
    return f'{int(m.group(1)):02d}-{int(m.group(2)):02d}' if m else str(s)

C_DARK='FF1B3A5C'; C_MED='FF2E6DA4'; C_LIGHT='FF4BACC6'
C_WHITE='FFFFFFFF'; C_STRIPE='FFDDEEFF'; C_BLACK='FF1A1A2E'
NO_BORDER = Border()

def mf(c=C_BLACK, b=False, s=11): return Font(name='Arial', color=c, bold=b, size=s)
def mfill(h): return PatternFill('solid', fgColor=h)
def mal(h='center', v='center', w=False): return Alignment(horizontal=h, vertical=v, wrap_text=w)

# ============================================================
# STEP 1: ตรวจสอบไฟล์
# ============================================================
print("\n📂 โหลดไฟล์ข้อมูล...")
for path, label in [
    (FILE_DAILY,'ยอดขายรายวัน'), (FILE_CATEGORY,'หมวดหมู่'),
    (FILE_PRODUCTS,'สินค้า'), (FILE_HOURLY,'ช่วงเวลา'),
    (FILE_RECEIPTS,'ใบเสร็จ'), (TEMPLATE_FILE,'Template'),
]:
    if not os.path.exists(path):
        print(f"❌ ไม่พบไฟล์ {label}: {path}"); sys.exit(1)
    print(f"   ✅ {label}")

# ============================================================
# STEP 2: คลีน data
# ============================================================

# --- Daily ---
daily_df = pd.read_excel(FILE_DAILY)
daily_df = daily_df[daily_df['วันที่'].astype(str).str.contains('/')].copy()
daily_df['sales']  = daily_df['ยอดขายรวม (฿)'].apply(parse_thb)
daily_df['net']    = daily_df['ยอดขายสุทธิ (฿)'].apply(parse_thb)
daily_df['orders'] = pd.to_numeric(daily_df['จำนวนออเดอร์'], errors='coerce').fillna(0).astype(int)
daily_df['disc']   = daily_df['ส่วนลด (฿)'].apply(parse_thb)
day_th = {0:'จันทร์',1:'อังคาร',2:'พุธ',3:'พฤหัส',4:'ศุกร์',5:'เสาร์',6:'อาทิตย์'}
daily_df['dt']       = pd.to_datetime(daily_df['วันที่'], dayfirst=True)
daily_df['day_th']   = daily_df['dt'].dt.dayofweek.map(day_th)
daily_df['date_str'] = daily_df['dt'].dt.strftime('%d/%m')
daily_df = daily_df.reset_index(drop=True)

# --- Category ---
cat_df = pd.read_excel(FILE_CATEGORY)
cat_df = cat_df[~cat_df['ชื่อหมวดหมู่'].isin(['รวมทั้งหมด','temp'])].copy()
cat_df['sales'] = cat_df['ยอดขายรวม (฿)'].apply(parse_thb)
cat_df['qty']   = pd.to_numeric(cat_df['จำนวน'], errors='coerce').fillna(0).astype(int)
cat_df['pct']   = (cat_df['sales'] / cat_df['sales'].sum() * 100).round(1)
cat_df = cat_df.sort_values('sales', ascending=False).reset_index(drop=True)

# --- Products ---
prod_df = pd.read_excel(FILE_PRODUCTS)
prod_df = prod_df[prod_df['ชื่อสินค้า'] != 'รวมทั้งหมด'].copy()
prod_df['sales'] = prod_df['ยอดขายสุทธิ (฿)'].apply(parse_thb)
top10 = prod_df.sort_values('sales', ascending=False).head(10).reset_index(drop=True)

# --- Hourly ---
hour_df = pd.read_excel(FILE_HOURLY)
hour_df = hour_df[hour_df['ระยะเวลา'] != 'รวมทั้งหมด'].copy()
hour_df['sales']  = hour_df['ยอดขายรวม (฿)'].apply(parse_thb)
hour_df['orders'] = pd.to_numeric(hour_df['จำนวนออเดอร์'], errors='coerce').fillna(0).astype(int)
hour_df = hour_df[hour_df['sales'] > 0].copy()
hour_df['time_str']   = hour_df['ระยะเวลา'].apply(fmt_hour)
hour_df['hour_start'] = hour_df['ระยะเวลา'].str.extract(r'^(\d+)').astype(int)

# เรียงใหม่: เย็น-ดึก (18:00+) ก่อน แล้วตามด้วยเช้า (00:00-05:59)
hour_evening = hour_df[hour_df['hour_start'] >= 18].sort_values('hour_start')
hour_night   = hour_df[hour_df['hour_start'] <  6].sort_values('hour_start')
hour_df = pd.concat([hour_evening, hour_night]).reset_index(drop=True)

# คำนวณ row positions จาก data จริง (ไม่ hardcode)
HOURLY_START = 7
HOURLY_ROWS  = len(hour_df)
HOURLY_END   = HOURLY_START + HOURLY_ROWS - 1
HOURLY_TOTAL = HOURLY_END + 1

print(f"   → Hourly: {HOURLY_ROWS} ช่วงเวลา (row {HOURLY_START}-{HOURLY_END}, total row {HOURLY_TOTAL})")

# --- Receipts ---
rec_df    = pd.read_excel(FILE_RECEIPTS)
num_bills = len(rec_df[rec_df['เลขที่ใบเสร็จ'].astype(str) != '-'])

# --- Summary Stats ---
total_sales  = daily_df['sales'].sum()
total_items  = int(rec_df['จำนวน'].sum())
total_disc   = daily_df['disc'].sum()
avg_per_bill = round(total_sales / num_bills) if num_bills > 0 else 0
n_open       = len(daily_df[daily_df['sales'] > 0])

print(f"\n📊 สรุปสถิติ")
print(f"   ยอดขายรวม : {total_sales:,.0f} บาท")
print(f"   บิล       : {num_bills}  |  เปิด {n_open}/{TOTAL_DAYS} วัน")
print(f"   รายการ    : {total_items:,}  |  เฉลี่ย/บิล {avg_per_bill:,} บาท")
print(f"   ส่วนลด   : {total_disc:,.0f} บาท ({round(total_disc/total_sales*100,1) if total_sales else 0:.1f}%)")

# ============================================================
# STEP 3: สร้าง Excel
# ============================================================
print("\n🏗️  สร้าง Dashboard...")

shutil.copy(TEMPLATE_FILE, OUTPUT_FILE)
wb = load_workbook(OUTPUT_FILE)
ws = wb.active

# Clear ทุกอย่าง
for mr in list(ws.merged_cells.ranges): ws.unmerge_cells(str(mr))
for row in ws.iter_rows():
    for cell in row:
        cell.value         = None
        cell.fill          = PatternFill('none')
        cell.border        = NO_BORDER
        cell.number_format = 'General'

# ── Title ──────────────────────────────────────────────────
ws['B1'].value = f'🏪  SALES DASHBOARD  |  ยอดขายประจำเดือน {MONTH_LABEL}  ({DATE_RANGE})'
ws['B1'].fill = mfill(C_DARK); ws['B1'].font = mf(C_WHITE, True, 14)
ws['B1'].alignment = mal(); ws.merge_cells('B1:J1')
ws.row_dimensions[1].height = 45

# ── KPI Cards row 2-3 ──────────────────────────────────────
for coord, txt, fill in [
    ('B2', f'💰 ยอดขายรวม\n{total_sales:,.0f} บาท',  C_MED),
    ('D2', f'📦 รายการขาย\n{total_items:,} รายการ',   C_LIGHT),
    ('H2', f' เฉลี่ย/บิล\n~{avg_per_bill:,} บาท',    C_DARK),
    ('J2', f'🏷️ ส่วนลด\n{total_disc:,.0f} บาท',       C_MED),
]:
    ws[coord].value = txt; ws[coord].fill = mfill(fill)
    ws[coord].font = mf(C_WHITE, True, 11); ws[coord].alignment = mal(w=True)
ws.merge_cells('B2:C3'); ws.merge_cells('D2:F3')
ws.merge_cells('H2:I3'); ws.merge_cells('J2:J3')
ws.row_dimensions[2].height = 49.5; ws.row_dimensions[3].height = 49.5

# ── Section headers row 5 ──────────────────────────────────
ws['B5'].value = '📅  ยอดขายรายวัน'
ws['B5'].fill = mfill(C_MED); ws['B5'].font = mf(C_WHITE, True)
ws['B5'].alignment = mal(); ws.merge_cells('B5:F5')
ws.row_dimensions[5].height = 25.5
ws['H5'].value = '⏰  ยอดขายตามช่วงเวลา'
ws['H5'].fill = mfill(C_MED); ws['H5'].font = mf(C_WHITE, True)
ws['H5'].alignment = mal(); ws.merge_cells('H5:J5')

# ── Column headers row 6 ───────────────────────────────────
for col, lbl in [('B','วันที่'),('C','วัน'),('D','ยอดขาย'),('E','ออเดอ'),('F','ยอดสุทธิ')]:
    c = ws[f'{col}6']; c.value = lbl
    c.fill = mfill(C_DARK); c.font = mf(C_WHITE, True); c.alignment = mal()
for col, lbl in [('H','ช่วงเวลา'),('I','ยอดขาย'),('J','ออเดอ')]:
    c = ws[f'{col}6']; c.value = lbl
    c.fill = mfill(C_DARK); c.font = mf(C_WHITE, True); c.alignment = mal()
ws.row_dimensions[6].height = 19.5

for col, w in [('A',2),('B',20),('C',14),('D',14),('E',8),('F',14),('G',3),('H',22),('I',14),('J',10)]:
    ws.column_dimensions[col].width = w

# ── Daily data rows 7-37 ───────────────────────────────────
for day_idx in range(TOTAL_DAYS):
    r = 7 + day_idx
    fill = C_STRIPE if day_idx % 2 == 0 else C_WHITE
    ws.row_dimensions[r].height = 19.5
    if day_idx < len(daily_df):
        row = daily_df.iloc[day_idx]
        vals = [('B', row['date_str']), ('C', row['day_th']),
                ('D', round(row['sales'])), ('E', int(row['orders'])), ('F', round(row['net']))]
    else:
        vals = [('B',''),('C',''),('D',0),('E',0),('F',0)]
    for col, val in vals:
        c = ws[f'{col}{r}']; c.value = val
        c.fill = mfill(fill); c.font = mf(C_BLACK); c.alignment = mal()
        c.border = NO_BORDER

# ── Hourly data (เรียง เย็น→ดึก→เช้า, ไม่มีเส้น) ──────────
for i, row in hour_df.iterrows():
    r = HOURLY_START + i
    fill = C_STRIPE if i % 2 == 0 else C_WHITE
    for col, val in [('H', row['time_str']), ('I', int(row['sales'])), ('J', int(row['orders']))]:
        c = ws[f'{col}{r}']; c.value = val
        c.fill = mfill(fill); c.font = mf(C_BLACK); c.alignment = mal()
        c.border = NO_BORDER

# ── Hourly total (คำนวณ row อัตโนมัติ) ────────────────────
ws[f'H{HOURLY_TOTAL}'].value = 'รวม'
ws[f'H{HOURLY_TOTAL}'].fill = mfill(C_MED); ws[f'H{HOURLY_TOTAL}'].font = mf(C_WHITE, True)
ws[f'H{HOURLY_TOTAL}'].alignment = mal()
ws[f'I{HOURLY_TOTAL}'].value = f'=SUM(I{HOURLY_START}:I{HOURLY_END})'
ws[f'I{HOURLY_TOTAL}'].fill = mfill(C_MED); ws[f'I{HOURLY_TOTAL}'].font = mf(C_WHITE, True)
ws[f'I{HOURLY_TOTAL}'].alignment = mal()
ws[f'J{HOURLY_TOTAL}'].value = f'=SUM(J{HOURLY_START}:J{HOURLY_END})'
ws[f'J{HOURLY_TOTAL}'].fill = mfill(C_MED); ws[f'J{HOURLY_TOTAL}'].font = mf(C_WHITE, True)
ws[f'J{HOURLY_TOTAL}'].alignment = mal()
ws.row_dimensions[HOURLY_TOTAL].height = 19.5

# ── Daily total row 38 ─────────────────────────────────────
ws['B38'].value = 'รวม'
ws['B38'].fill = mfill(C_MED); ws['B38'].font = mf(C_WHITE, True); ws['B38'].alignment = mal()
ws['C38'].value = n_open
ws['C38'].fill = mfill(C_MED); ws['C38'].font = mf(C_WHITE, True); ws['C38'].alignment = mal()
for col in ['D','E','F']:
    c = ws[f'{col}38']; c.value = f'=SUM({col}7:{col}37)'
    c.fill = mfill(C_MED); c.font = mf(C_WHITE, True); c.alignment = mal()
ws.row_dimensions[38].height = 19.5

# ── Gap row 39 ─────────────────────────────────────────────
ws.row_dimensions[39].height = 7.5

# ── Section headers row 40 ─────────────────────────────────
ws['B40'].value = '🗂️  ยอดขายตามหมวดหมู่'
ws['B40'].fill = mfill(C_MED); ws['B40'].font = mf(C_WHITE, True); ws['B40'].alignment = mal()
ws.merge_cells('B40:E40')
ws['H40'].value = '🏆  TOP 10 สินค้าขายดี'
ws['H40'].fill = mfill(C_MED); ws['H40'].font = mf(C_WHITE, True); ws['H40'].alignment = mal()
ws.merge_cells('H40:J40')
ws.row_dimensions[40].height = 19.5

# ── Column headers row 41 ──────────────────────────────────
for col, lbl in [('B','หมวดหมู่'),('C','ยอดขาย'),('D','จำนวน'),('E','%')]:
    c = ws[f'{col}41']; c.value = lbl
    c.fill = mfill(C_DARK); c.font = mf(C_WHITE, True); c.alignment = mal()
for col, lbl in [('H','สินค้า'),('I','ยอดขาย'),('J','จำนวน')]:
    c = ws[f'{col}41']; c.value = lbl
    c.fill = mfill(C_DARK); c.font = mf(C_WHITE, True); c.alignment = mal()
ws.row_dimensions[41].height = 19.5

# ── Category data rows 42-50 ───────────────────────────────
for i in range(9):
    r = 42 + i; fill = C_STRIPE if i % 2 == 0 else C_WHITE
    ws.row_dimensions[r].height = 19.5
    if i < len(cat_df):
        row = cat_df.iloc[i]
        for col, val in [('B', row['ชื่อหมวดหมู่']), ('C', int(row['sales'])),
                         ('D', row['qty']), ('E', row['pct'])]:
            c = ws[f'{col}{r}']; c.value = val
            c.fill = mfill(fill); c.font = mf(C_BLACK); c.alignment = mal()
    else:
        for col in ['B','C','D','E']:
            c = ws[f'{col}{r}']; c.value = ''; c.fill = mfill(fill)

# ── Top10 data rows 42-51 ──────────────────────────────────
for i in range(10):
    r = 42 + i; fill = C_STRIPE if i % 2 == 0 else C_WHITE
    ws.row_dimensions[r].height = 19.5
    if i < len(top10):
        row = top10.iloc[i]
        for col, val in [('H', f'{i+1}. {row["ชื่อสินค้า"]}'),
                         ('I', int(row['sales'])), ('J', int(row['จำนวน']))]:
            c = ws[f'{col}{r}']; c.value = val
            c.fill = mfill(fill); c.font = mf(C_BLACK)
            c.alignment = mal(h='left' if col == 'H' else 'center')

# ── Category total row 51 ──────────────────────────────────
ws['B51'].value = 'รวม'
ws['B51'].fill = mfill(C_MED); ws['B51'].font = mf(C_WHITE, True); ws['B51'].alignment = mal()
for col in ['C','D','E']:
    c = ws[f'{col}51']; c.value = f'=SUM({col}42:{col}50)'
    c.fill = mfill(C_MED); c.font = mf(C_WHITE, True); c.alignment = mal()
ws.row_dimensions[51].height = 19.5

# ============================================================
# STEP 4: Save
# ============================================================
wb.save(OUTPUT_FILE)
print(f"\n✅ บันทึกสำเร็จ → {OUTPUT_FILE}")
print("=" * 50)