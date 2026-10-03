import argparse
import sys
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

WEEKDAYS = [
    ("จันทร์", "จ"), ("อังคาร", "อ"), ("พุธ", "พ"), ("พฤหัส", "พฤ"),
    ("ศุกร์", "ศ"), ("เสาร์", "ส"), ("อาทิตย์", "อา"),
]
ABBRS = {abbr for _, abbr in WEEKDAYS}
MONTH_ORDER = ["ม.ค.", "ก.พ.", "มี.ค.", "เม.ย.", "พ.ค.", "มิ.ย.",
               "ก.ค.", "ส.ค.", "ก.ย.", "ต.ค.", "พ.ย.", "ธ.ค."]
FONT_NAME = "Arial"


# ---------------------------------------------------------------- อ่านข้อมูล
def _find_columns(header_row_values, sheet):
    """หาตำแหน่งคอลัมน์จาก "หัวตาราง" (แถว 5) แทนการฟิกซ์ตัวอักษรคอลัมน์
    - วัน        : เซลล์ที่เป็น 'วัน' เป๊ะ ๆ
    - รายรับรวม  : 'Total' ตัวแรก (กลุ่มรายรับ FOOD/DRINK/Total)
    - กำไรสุทธิ  : เซลล์ที่มีคำว่า 'กำไรสุทธิ'  (ชีตบางเดือนอยู่ Y บางเดือนอยู่ Z)
    """
    day = revenue = profit = None
    for i, v in enumerate(header_row_values):
        t = str(v).strip() if v is not None else ""
        if day is None and t == "วัน":
            day = i
        if revenue is None and t == "Total":
            revenue = i
        if profit is None and "กำไรสุทธิ" in t:
            profit = i
    missing = [n for n, x in (("วัน", day), ("รายรับ Total", revenue),
                              ("กำไรสุทธิ", profit)) if x is None]
    if missing:
        raise ValueError(f"ชีต '{sheet}': หาหัวคอลัมน์ไม่เจอ: {', '.join(missing)} "
                         f"(ตรวจสอบว่าแถวหัวตารางอยู่แถวที่ระบุใน --header-row)")
    return day, revenue, profit


def _num(v):
    """ตัวเลข -> float, อย่างอื่น ('-', ว่าง, ข้อความ) -> NaN"""
    if isinstance(v, bool):
        return float("nan")
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        return float("nan")


def _load_daily_data(wb, sheet, header_row, first_row, last_row, skip_unfilled):
    """อ่านข้อมูลรายวันของ 1 ชีต -> DataFrame: day, revenue, profit"""
    ws = wb[sheet]
    rows = list(ws.iter_rows(min_row=header_row, max_row=last_row, values_only=True))
    header = rows[0]
    data = rows[first_row - header_row:]
    day_i, rev_i, prof_i = _find_columns(header, sheet)

    def cell(row, i):
        return row[i] if i < len(row) else None

    recs = []
    for row in data:
        d = cell(row, day_i)
        d = str(d).strip() if d is not None else ""
        if d not in ABBRS:          # ข้ามแถว 'รวมเดือน', OVERHEAD, แถวว่าง ฯลฯ
            continue
        recs.append({"day": d,
                     "revenue": _num(cell(row, rev_i)),
                     "profit": _num(cell(row, prof_i))})
    df = pd.DataFrame(recs, columns=["day", "revenue", "profit"])

    if skip_unfilled and df["revenue"].notna().any():
        # ตัดแถวท้ายเดือนที่ยังไม่ได้กรอกออก (ไม่นับเป็น "วันปิด")
        last = df["revenue"].last_valid_index()
        df = df.loc[:last]
    elif skip_unfilled:
        df = df.iloc[0:0]
    return df


def resolve_sheets(arg, available):
    """แปลงค่า --sheets เป็นรายชื่อชีต
    รองรับ:  "all" | "ม.ค.,ก.พ.,มี.ค." | "ม.ค.-มี.ค." | ผสมกันได้ เช่น "ม.ค.-มี.ค.,ก.ค."
    """
    avail = {name.strip(): name for name in available}   # กันช่องว่างหน้า/หลังชื่อชีต
    if arg.strip().lower() == "all":
        return [avail[m] for m in MONTH_ORDER if m in avail]

    out = []
    for tok in arg.split(","):
        tok = tok.strip()
        if not tok:
            continue
        if tok in avail:
            out.append(avail[tok])
        elif "-" in tok:
            a, b = [x.strip() for x in tok.split("-", 1)]
            if a not in MONTH_ORDER or b not in MONTH_ORDER:
                raise ValueError(f"อ่านช่วงเดือน '{tok}' ไม่ได้ (ใช้ชื่อย่อ เช่น ม.ค.-มี.ค.)")
            ia, ib = MONTH_ORDER.index(a), MONTH_ORDER.index(b)
            if ia > ib:
                raise ValueError(f"ช่วงเดือน '{tok}' เรียงกลับด้าน")
            for m in MONTH_ORDER[ia:ib + 1]:
                if m not in avail:
                    raise ValueError(f"ไม่พบชีต '{m}' ในไฟล์")
                out.append(avail[m])
        else:
            raise ValueError(f"ไม่พบชีต '{tok}' ในไฟล์\nชีตที่มี: {list(avail)}")
    seen, uniq = set(), []
    for s in out:                    # ตัดชีตซ้ำ แต่คงลำดับ
        if s not in seen:
            seen.add(s)
            uniq.append(s)
    if not uniq:
        raise ValueError("ไม่ได้เลือกชีตใดเลย")
    return uniq


# ---------------------------------------------------------------- คำนวณสรุป
def compute_weekday_summary(path, month_sheets, header_row=5, data_start_row=6,
                            data_end_row=40, skip_unfilled=False):
    """รวมข้อมูลทุกชีตที่เลือก แล้วสรุปตามวันในสัปดาห์ (จ.-อา.) เป็นค่าจริง"""
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        sheets = resolve_sheets(",".join(month_sheets), wb.sheetnames)
        frames = [_load_daily_data(wb, s, header_row, data_start_row,
                                   data_end_row, skip_unfilled) for s in sheets]
    finally:
        wb.close()
    all_days = pd.concat(frames, ignore_index=True)

    rows = []
    for th_name, abbr in WEEKDAYS:
        sub = all_days[all_days["day"] == abbr]
        open_days = int(sub["revenue"].notna().sum())
        closed_days = int(sub["revenue"].isna().sum())
        total_rev = sub["revenue"].sum()
        total_profit = sub["profit"].sum()
        avg_rev = total_rev / open_days if open_days else 0
        avg_profit = total_profit / open_days if open_days else 0
        rows.append({
            "วัน": th_name, "ตัวย่อ": abbr,
            "จำนวนวันเปิด": open_days, "จำนวนวันปิด": closed_days,
            "รายรับเฉลี่ย/วัน (฿)": round(avg_rev),
            "กำไรเฉลี่ย/วัน (฿)": round(avg_profit),
            "รายรับรวม (฿)": round(total_rev),
            "กำไรรวม (฿)": round(total_profit),
        })
    return pd.DataFrame(rows), sheets


# ---------------------------------------------------------------- เขียนไฟล์
def _fill_sheet(ws, summary):
    header_font = Font(name=FONT_NAME, size=12, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2A78D6")
    normal_font = Font(name=FONT_NAME, size=11)
    bold_font = Font(name=FONT_NAME, size=11, bold=True)
    side = Side(style="thin", color="CCCCCC")
    border = Border(left=side, right=side, top=side, bottom=side)
    center = Alignment(horizontal="center", vertical="center")
    right = Alignment(horizontal="right")

    headers = [h.replace(" (฿)", "") for h in summary.columns]
    for i, h in enumerate(headers):
        c = ws.cell(row=1, column=1 + i, value=h)
        c.font, c.fill, c.alignment, c.border = header_font, header_fill, center, border

    for r_off, (_, row) in enumerate(summary.iterrows()):
        for c_off, col in enumerate(summary.columns):
            val = row[col]
            if hasattr(val, "item"):          # numpy -> python
                val = val.item()
            c = ws.cell(row=2 + r_off, column=1 + c_off, value=val)
            c.border = border
            if col in ("วัน", "ตัวย่อ", "จำนวนวันเปิด", "จำนวนวันปิด"):
                c.font, c.alignment = normal_font, center
            else:
                c.font = bold_font if "รวม" in col else normal_font
                c.alignment = right
                c.number_format = "#,##0"

    for col, w in {"A": 12, "B": 10, "C": 14, "D": 14,
                   "E": 18, "F": 18, "G": 16, "H": 16}.items():
        ws.column_dimensions[col].width = w


def write_summary(output_path, summary, output_sheet, source_path=None):
    """source_path=None  -> สร้างไฟล์ใหม่ที่มีแค่ชีตสรุป (ปลอดภัยสุด)
       source_path=...   -> คัดลอกทั้งไฟล์มาแล้วเพิ่มชีตสรุปไว้หน้าสุด
    ไม่ว่าแบบไหน ไฟล์ต้นฉบับไม่ถูกแก้"""
    if source_path:
        wb = load_workbook(source_path)
        if output_sheet in wb.sheetnames:
            del wb[output_sheet]
        ws = wb.create_sheet(output_sheet, 0)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = output_sheet
    _fill_sheet(ws, summary)
    wb.save(output_path)


# ---------------------------------------------------------------- CLI
def _cli():
    p = argparse.ArgumentParser(description="สรุปยอดขาย/กำไรตามวันในสัปดาห์ จากชีตรายเดือน")
    p.add_argument("path", help="ไฟล์ต้นฉบับ (จะไม่ถูกแก้ไข)")
    p.add_argument("--sheets", required=True,
                   help='ชีตที่เลือก: "ม.ค.,ก.พ.,มี.ค." หรือ "ม.ค.-มี.ค." หรือ "all"')
    p.add_argument("--output", default=None,
                   help="ไฟล์ผลลัพธ์ (default: <ชื่อเดิม>_summary.xlsx)")
    p.add_argument("--output-sheet", default="สรุปวันขายดี")
    p.add_argument("--into-copy", action="store_true",
                   help="สร้างเป็นสำเนาของทั้งไฟล์ + เพิ่มชีตสรุปไว้หน้าสุด "
                        "(ไม่ใส่ = ไฟล์ใหม่ที่มีเฉพาะชีตสรุป)")
    p.add_argument("--skip-unfilled", action="store_true",
                   help="ไม่นับวันท้ายเดือนที่ยังไม่ได้กรอก (ช่อง '-' หลังวันสุดท้ายที่มีข้อมูล) "
                        "ว่าเป็นวันปิด -- ใช้กับเดือนปัจจุบันที่ยังไม่จบ")
    p.add_argument("--header-row", type=int, default=5)
    p.add_argument("--data-start-row", type=int, default=6)
    p.add_argument("--data-end-row", type=int, default=40,
                   help="แถวสุดท้ายที่จะสแกน (แถวที่ไม่ใช่วันจะถูกข้ามเอง)")
    args = p.parse_args()

    output_path = args.output
    if output_path is None:
        stem, ext = args.path.rsplit(".", 1)
        output_path = f"{stem}_summary.{ext}"

    try:
        summary, sheets = compute_weekday_summary(
            args.path, [args.sheets],
            header_row=args.header_row, data_start_row=args.data_start_row,
            data_end_row=args.data_end_row, skip_unfilled=args.skip_unfilled)
        write_summary(output_path, summary, args.output_sheet,
                      source_path=args.path if args.into_copy else None)
    except (ValueError, FileNotFoundError, PermissionError) as e:
        sys.exit(f"ERROR: {e}")

    print(f"ชีตที่ใช้ ({len(sheets)}): {', '.join(sheets)}\n")
    print(summary.to_string(index=False))
    print(f"\nDone. ไฟล์ต้นฉบับ '{args.path}' ไม่ถูกแก้ไข -> สร้างไฟล์ใหม่ที่ '{output_path}'")


if __name__ == "__main__":
    _cli()